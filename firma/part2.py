"""LangGraph-based source profiling, mapping, validation, and extraction."""

from __future__ import annotations

import csv
import io
import json
import os
import re
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Literal, NotRequired, TypedDict
from urllib.parse import urlparse
import xml.etree.ElementTree as ET

import requests
import yaml
from dotenv import load_dotenv
from langgraph.graph import END, START, StateGraph
from pydantic import BaseModel, Field

ENGINE_VERSION = "part2-v2"
SAMPLE_LIMIT = 1200
DOWNLOAD_LIMIT = 32 * 1024 * 1024
MAX_LLM_REVISIONS = 1
SUPPORTED_FORMATS = {"csv", "tsv", "json", "jsonl", "xml", "xlsx", "xls", "zip"}


class ProposedField(BaseModel):
    source_field: str
    target_field: str
    confidence: float = Field(ge=0.0, le=1.0)
    rationale: str
    transformation: Literal[
        "identity", "normalize_whitespace", "digits_only", "date_iso", "status_enum"
    ]


class MappingProposal(BaseModel):
    fields: list[ProposedField]
    notes: list[str]


class AgentState(TypedDict):
    source: dict[str, Any]
    ontology: dict[str, Any]
    sample_limit: NotRequired[int]
    profile: NotRequired[dict[str, Any]]
    records: NotRequired[list[dict[str, Any]]]
    mapping: NotRequired[dict[str, Any]]
    validation: NotRequired[dict[str, Any]]
    revisions: NotRequired[list[dict[str, Any]]]
    revision_count: NotRequired[int]
    llm_model: NotRequired[str]
    observations: NotRequired[list[dict[str, Any]]]
    error: NotRequired[str]
    approval: NotRequired[str]


def _slug(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", value.lower()).strip()


def _format(resource: dict[str, Any]) -> str:
    value = str(resource.get("format") or "").lower().replace("excel", "xlsx")
    if "spreadsheet" in value:
        return "xlsx"
    if "comma" in value:
        return "csv"
    if "tab" in value:
        return "tsv"
    if "json" in value:
        return "json"
    if "xml" in value:
        return "xml"
    if "zip" in value:
        return "zip"
    if "csv" in value:
        return "csv"
    if "tsv" in value:
        return "tsv"
    ext = Path(urlparse(str(resource.get("url", ""))).path).suffix.lower().lstrip(".")
    return "xlsx" if ext == "xls" else ext


def select_sources(shortlist_path: Path, limit: int = 6) -> list[dict[str, Any]]:
    """Choose downloadable structured resources from Part 1 metadata, with format diversity."""
    datasets = [json.loads(line) for line in shortlist_path.read_text(encoding="utf-8").splitlines() if line.strip()]
    options: list[tuple[int, dict[str, Any], dict[str, Any]]] = []
    for dataset in datasets:
        for resource in dataset.get("resources", []):
            fmt = _format(resource)
            name = _slug(str(resource.get("name", "")))
            if fmt not in SUPPORTED_FORMATS or not resource.get("url"):
                continue
            url = str(resource.get("url", "")).lower()
            if any(word in name for word in ("schema", "readme", "help file", "documentation", "data dictionary", "resource list")):
                continue
            if "?wsdl" in url or url.endswith(".pdf"):
                continue
            if resource.get("size") and int(resource["size"]) > DOWNLOAD_LIMIT and not resource.get("datastore_active"):
                continue
            rank = {"csv": 0, "tsv": 1, "json": 2, "jsonl": 3, "xml": 4, "xlsx": 5, "xls": 6, "zip": 7}.get(fmt, 8)
            if resource.get("datastore_active"):
                rank -= 3
            options.append((rank, dataset, resource))
    options.sort(key=lambda item: (item[0], -float(item[1].get("confidence", 0))))

    selected: list[dict[str, Any]] = []
    used_dataset_ids: set[str] = set()
    used_formats: set[str] = set()
    for diverse_pass in (True, False):
        for _, dataset, resource in options:
            source_id = str(dataset.get("id", ""))
            fmt = _format(resource)
            if source_id in used_dataset_ids or (diverse_pass and fmt in used_formats):
                continue
            selected.append({
                "source_id": source_id,
                "title": dataset.get("title", "Untitled"),
                "organization": (dataset.get("organization") or {}).get("title", "unknown"),
                "licence": dataset.get("license_title") or dataset.get("license_id") or "unknown",
                "resource": resource,
                "format": fmt,
                "portal": resource.get("ckan_url") or "https://data.gov.au",
            })
            used_dataset_ids.add(source_id)
            used_formats.add(fmt)
            if len(selected) == limit:
                return selected
    return selected


def _ckan_records(source: dict[str, Any], limit: int) -> list[dict[str, Any]] | None:
    resource = source["resource"]
    if not resource.get("datastore_active") or not resource.get("id"):
        return None
    portal = str(source.get("portal", "https://data.gov.au")).rstrip("/")
    parsed = urlparse(portal)
    api_root = portal if "/data" in parsed.path else f"{parsed.scheme}://{parsed.netloc}/data"
    response = requests.get(f"{api_root}/api/3/action/datastore_search",
                            params={"resource_id": resource["id"], "limit": limit}, timeout=45)
    response.raise_for_status()
    payload = response.json()
    if not payload.get("success"):
        raise ValueError(f"CKAN datastore_search failed for {resource['id']}")
    return [row for row in payload.get("result", {}).get("records", []) if isinstance(row, dict)]


def _read_direct(source: dict[str, Any], limit: int) -> list[dict[str, Any]]:
    resource = source["resource"]
    response = requests.get(str(resource["url"]), stream=True, timeout=60)
    response.raise_for_status()
    content = bytearray()
    for chunk in response.iter_content(64 * 1024):
        content.extend(chunk)
        if len(content) >= DOWNLOAD_LIMIT:
            break
    response.close()
    raw, fmt = bytes(content), source["format"]
    if fmt in {"csv", "tsv"}:
        text = raw.decode("utf-8-sig", errors="replace")
        dialect = csv.Sniffer().sniff(text[:65536], delimiters=",;\t|")
        return [dict(row) for _, row in zip(range(limit), csv.DictReader(io.StringIO(text), dialect=dialect))]
    if fmt in {"json", "jsonl"}:
        text = raw.decode("utf-8-sig", errors="replace")
        try:
            payload = json.loads(text)
            if isinstance(payload, list):
                return [item for item in payload[:limit] if isinstance(item, dict)]
            if isinstance(payload, dict) and isinstance(payload.get("features"), list):
                return [{**(feature.get("properties") or {}), "geometry": feature.get("geometry")}
                        for feature in payload["features"][:limit] if isinstance(feature, dict)]
            for key in ("records", "results", "data", "items"):
                if isinstance(payload, dict) and isinstance(payload.get(key), list):
                    return [item for item in payload[key][:limit] if isinstance(item, dict)]
        except json.JSONDecodeError:
            return [json.loads(line) for line in text.splitlines()[:limit] if line.strip()]
        raise ValueError("JSON resource has no recognizable record list")
    if fmt == "xml":
        root = ET.fromstring(raw)
        return [{child.tag.rsplit("}", 1)[-1]: child.text for child in row} for row in list(root)[:limit]]
    if fmt in {"xlsx", "xls"}:
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        rows = workbook.active.iter_rows(values_only=True)
        headers = [str(value or f"column_{index}") for index, value in enumerate(next(rows), 1)]
        return [dict(zip(headers, row)) for row in rows if any(value is not None for value in row)][:limit]
    if fmt == "zip":
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            members = [name for name in archive.namelist() if Path(name).suffix.lower() in {".csv", ".tsv", ".json", ".jsonl", ".xml"}]
            if not members:
                raise ValueError("ZIP contains no supported tabular or XML member")
            member = sorted(members)[0]
            nested = archive.read(member)[:DOWNLOAD_LIMIT]
            suffix = Path(member).suffix.lower()
            if suffix in {".csv", ".tsv"}:
                text = nested.decode("utf-8-sig", errors="replace")
                dialect = csv.Sniffer().sniff(text[:65536], delimiters=",;\t|")
                return [dict(row) for _, row in zip(range(limit), csv.DictReader(io.StringIO(text), dialect=dialect))]
            if suffix in {".json", ".jsonl"}:
                try:
                    payload = json.loads(nested)
                    return [row for row in payload[:limit] if isinstance(row, dict)] if isinstance(payload, list) else []
                except json.JSONDecodeError:
                    return [json.loads(line) for line in nested.decode("utf-8-sig").splitlines()[:limit] if line.strip()]
            root = ET.fromstring(nested)
            return [{child.tag.rsplit("}", 1)[-1]: child.text for child in row} for row in list(root)[:limit]]
    raise ValueError(f"Unsupported source format: {fmt}")


def load_sample(source: dict[str, Any], limit: int = SAMPLE_LIMIT) -> list[dict[str, Any]]:
    records = _ckan_records(source, limit)
    return records if records is not None else _read_direct(source, limit)


def _flatten_keys(record: dict[str, Any], prefix: str = "") -> dict[str, Any]:
    flattened: dict[str, Any] = {}
    for key, value in record.items():
        name = f"{prefix}.{key}" if prefix else str(key)
        if isinstance(value, dict):
            flattened.update(_flatten_keys(value, name))
        else:
            flattened[name] = value
    return flattened


def _get_mapping_model():
    load_dotenv(Path.cwd() / ".env", override=False)
    if not os.getenv("OPENAI_API_KEY"):
        raise RuntimeError("OPENAI_API_KEY is missing. Set it in .env or the environment before running Part 2.")
    from langchain_openai import ChatOpenAI

    model_name = os.getenv("OPENAI_MODEL", "gpt-4o-mini")
    return ChatOpenAI(model=model_name, temperature=0).with_structured_output(MappingProposal), model_name


def _profile_prompt(profile: dict[str, Any], ontology: dict[str, Any]) -> str:
    compact_profile = {
        "source_title": profile.get("title"),
        "format": profile.get("format"),
        "sample_record_count": profile.get("sample_record_count"),
        "fields": profile.get("fields", [])[:160],
        "irregularities": profile.get("irregularities", []),
    }
    contract = {section: values for section, values in ontology.items() if section in {"entity", "address"}}
    return json.dumps({"profile": compact_profile, "ontology": contract}, ensure_ascii=False)


def _invoke_mapping_model(state: AgentState, revise: bool = False) -> tuple[MappingProposal, str]:
    structured_model, model_name = _get_mapping_model()
    if revise:
        instruction = (
            "You are revising a dataset-to-ontology mapping after deterministic validation. "
            "Return the complete revised mapping, not only changed fields. Use the validation failures "
            "and representative samples to correct a mapping or leave the source field unmapped. "
            "Never force a mapping to satisfy validation; abstain where ambiguous."
        )
        validation_context = {
            "previous_mapping": state["mapping"]["fields"],
            "validation": state["validation"],
            "revision_number": state.get("revision_count", 0) + 1,
        }
    else:
        instruction = (
            "Infer a mapping for this previously unseen public dataset. Use field names, sample values, "
            "types, layout, and ontology definitions together. Map only fields whose meaning is supported "
            "by evidence. Keep ambiguous fields unmapped. Do not infer entity identity or manufacture values. "
            "Use the narrowest valid transformation and provide a concise evidence-based rationale. "
            "The confidence is your uncalibrated mapping-confidence estimate, not a claim of truth."
        )
        validation_context = {}
    prompt = {
        "dataset_evidence": json.loads(_profile_prompt(state["profile"], state["ontology"])),
        "revision_context": validation_context,
        "output_rules": {
            "source_field": "must exactly match a profiled field name",
            "target_field": "must be a field defined under ontology.entity or ontology.address",
            "transformation_options": ["identity", "normalize_whitespace", "digits_only", "date_iso", "status_enum"],
            "confidence_minimum_for_proposal": 0.65,
        },
    }
    response = structured_model.invoke([
        ("system", instruction),
        ("human", json.dumps(prompt, ensure_ascii=False)),
    ])
    if not isinstance(response, MappingProposal):
        response = MappingProposal.model_validate(response)
    return response, model_name


def _validated_mapping(response: MappingProposal, state: AgentState) -> dict[str, Any]:
    profile_fields = {field["name"] for field in state["profile"].get("fields", [])}
    targets = {
        target: section
        for section in ("entity", "address")
        for target in state["ontology"].get(section, {})
    }
    accepted = []
    used_targets: set[str] = set()
    for proposal in response.fields:
        if proposal.source_field not in profile_fields or proposal.target_field not in targets:
            continue
        if proposal.confidence < 0.65 or proposal.target_field in used_targets:
            continue
        accepted.append({
            "source_field": proposal.source_field,
            "target_field": proposal.target_field,
            "section": targets[proposal.target_field],
            "confidence": proposal.confidence,
            "confidence_basis": "LLM self-assessment; uncalibrated; separate from held-out validation",
            "rationale": proposal.rationale,
            "transformation": proposal.transformation,
        })
        used_targets.add(proposal.target_field)
    mapped_fields = {field["source_field"] for field in accepted}
    return {
        "source_id": state["source"]["source_id"],
        "source_title": state["source"]["title"],
        "format": state["source"]["format"],
        "fields": accepted,
        "unmapped_source_fields": sorted(profile_fields - mapped_fields),
        "llm_notes": response.notes,
        "approval": "pending",
    }


def inspect_source(state: AgentState) -> dict[str, Any]:
    source = state["source"]
    try:
        records = load_sample(source, state.get("sample_limit", SAMPLE_LIMIT))
        if not records:
            raise ValueError("No records could be sampled")
        fields: dict[str, list[Any]] = {}
        for record in records:
            for key, value in _flatten_keys(record).items():
                fields.setdefault(key, []).append(value)
        profile = {
            "source_id": source["source_id"], "title": source["title"],
            "format": source["format"], "sample_record_count": len(records),
            "fields": [{"name": key, "non_null": sum(value not in (None, "") for value in values),
                        "examples": [str(value)[:120] for value in values if value not in (None, "")][:3],
                        "types": sorted({type(value).__name__ for value in values if value is not None})}
                       for key, values in fields.items()],
            "irregularities": ["Nested object paths flattened with dot notation"] if any("." in key for key in fields) else [],
        }
        return {"profile": profile, "records": records, "error": ""}
    except Exception as exc:
        return {"profile": {"source_id": source["source_id"], "format": source["format"], "fields": []}, "records": [], "error": str(exc)}


def _propose(state: AgentState) -> dict[str, Any]:
    response, model_name = _invoke_mapping_model(state)
    return {"mapping": _validated_mapping(response, state), "llm_model": model_name}


def _value_valid(target: str, value: Any) -> bool:
    if value is None or str(value).strip() == "":
        return False
    text = str(value).strip()
    if target == "abn":
        digits = re.sub(r"\D", "", text)
        if len(digits) != 11:
            return False
        weights = [10, 1, 3, 5, 7, 9, 11, 13, 15, 17, 19]
        return sum((int(digit) - (1 if index == 0 else 0)) * weights[index] for index, digit in enumerate(digits)) % 89 == 0
    if target in {"acn", "nzbn"}:
        return len(re.sub(r"\D", "", text)) == (9 if target == "acn" else 13)
    if target == "date_registered":
        try:
            datetime.fromisoformat(text.replace("Z", "+00:00"))
            return True
        except ValueError:
            return bool(re.fullmatch(r"\d{1,2}/\d{1,2}/\d{4}|\d{4}-\d{2}-\d{2}", text))
    return True


def _canonical_value(value: Any, transformation: str) -> str:
    text = " ".join(str(value).split())
    if transformation == "date_iso":
        for pattern in ("%d/%m/%Y", "%Y-%m-%d", "%Y-%m-%dT%H:%M:%S"):
            try:
                return datetime.strptime(text, pattern).date().isoformat()
            except ValueError:
                continue
    if transformation == "status_enum":
        normalized = text.lower().strip()
        if normalized in {"current", "active", "registered", "appr", "approved", "approved/current"}:
            return "active"
        if normalized in {"ceased", "cancelled", "canceled", "deregistered", "expired"}:
            return "deregistered"
        if "liquidation" in normalized:
            return "in_liquidation"
        if "suspend" in normalized:
            return "suspended"
        return "unknown"
    if transformation == "digits_only":
        return re.sub(r"\D", "", text)
    return text


def _validate(state: AgentState) -> dict[str, Any]:
    records = state["records"]
    heldout = records[max(1, len(records) * 2 // 3):] or records[-1:]
    checks = []
    for mapping in state["mapping"]["fields"]:
        values = [_flatten_keys(record).get(mapping["source_field"]) for record in heldout]
        nonempty = [value for value in values if value not in (None, "")]
        valid = sum(_value_valid(mapping["target_field"], value) for value in nonempty)
        ratio = valid / len(nonempty) if nonempty else 0.0
        checks.append({"source_field": mapping["source_field"], "target_field": mapping["target_field"],
                       "holdout_rows": len(heldout), "nonempty": len(nonempty), "valid": valid,
                       "valid_ratio": round(ratio, 4), "pass": ratio >= 0.8 and bool(nonempty)})
    failed = [item["source_field"] for item in checks if not item["pass"]]
    return {"validation": {"holdout_rows": len(heldout), "checks": checks, "failed_fields": failed,
                           "passed": bool(checks) and not failed, "step": "unseen holdout sample"}}


def _revise(state: AgentState) -> dict[str, Any]:
    response, model_name = _invoke_mapping_model(state, revise=True)
    revised_mapping = _validated_mapping(response, state)
    previous = state["mapping"]
    revised_sources = {field["source_field"] for field in revised_mapping["fields"]}
    return {
        "mapping": revised_mapping,
        "llm_model": model_name,
        "revision_count": state.get("revision_count", 0) + 1,
        "revisions": (state.get("revisions", []) + [{
            "revision": state.get("revision_count", 0) + 1,
            "previous_fields": previous["fields"],
            "revised_fields": revised_mapping["fields"],
            "removed_fields": [field["source_field"] for field in previous["fields"]
                               if field["source_field"] not in revised_sources],
            "validation_failures": state["validation"]["failed_fields"],
            "notes": revised_mapping["llm_notes"],
        }]),
    }


def _abstain(state: AgentState) -> dict[str, Any]:
    failed = set(state["validation"]["failed_fields"])
    mapping = dict(state["mapping"])
    mapping["fields"] = [field for field in mapping["fields"] if field["source_field"] not in failed]
    mapping["unmapped_source_fields"] = sorted(set(mapping["unmapped_source_fields"]) | failed)
    mapping["abstention_reason"] = "Fields still failed deterministic validation after one LLM revision."
    return {"mapping": mapping}


def _extract(state: AgentState) -> dict[str, Any]:
    source, mapping = state["source"], state["mapping"]
    ingested = datetime.now(timezone.utc).isoformat()
    observations = []
    for index, record in enumerate(state["records"]):
        flat = _flatten_keys(record)
        canonical: dict[str, Any] = {"entity": {}, "address": {}}
        provenance: dict[str, Any] = {}
        for field in mapping["fields"]:
            value = flat.get(field["source_field"])
            if value in (None, ""):
                continue
            normalized = _canonical_value(value, field["transformation"])
            if not _value_valid(field["target_field"], normalized):
                continue
            canonical[field["section"]][field["target_field"]] = normalized
            provenance[field["target_field"]] = {"source_field": field["source_field"],
                                                   "field_confidence": field["confidence"], "derivation_level": "L1"}
        record_id = next((flat[key] for key in flat if _slug(key.rsplit(".", 1)[-1]) in {"id", "record id", "company number"}), None)
        canonical["observation"] = {"source_id": source["source_id"],
                                    "source_record_id": str(record_id) if record_id not in (None, "") else f"{source['source_id']}:{index}",
                                    "observed_at": ingested, "ingested_at": ingested,
                                    "licence": source.get("licence", "unknown"), "extractor_version": ENGINE_VERSION}
        canonical["confidence"] = {"source_reliability": None, "field_confidence": provenance}
        if any(canonical[section] for section in ("entity", "address")):
            observations.append(canonical)
    return {"observations": observations, "approval": "pending"}


def build_agent():
    graph = StateGraph(AgentState)
    graph.add_node("inspect", inspect_source)
    graph.add_node("propose", _propose)
    graph.add_node("validate", _validate)
    graph.add_node("revise", _revise)
    graph.add_node("abstain", _abstain)
    graph.add_node("extract", _extract)
    graph.add_edge(START, "inspect")
    graph.add_conditional_edges("inspect", lambda state: "propose" if not state.get("error") else END)
    graph.add_edge("propose", "validate")
    graph.add_conditional_edges(
        "validate",
        lambda state: "extract" if state["validation"]["passed"] else (
            "revise" if state.get("revision_count", 0) < MAX_LLM_REVISIONS else "abstain"
        ),
    )
    graph.add_edge("revise", "validate")
    graph.add_edge("abstain", "extract")
    graph.add_edge("extract", END)
    return graph.compile()


def run_part2(shortlist_path: Path | str = "outputs/part1/shortlist.jsonl",
              ontology_path: Path | str = "assets/firma_ontology.yaml",
              output_dir: Path | str = "outputs/part2", limit: int = 6,
              sample_limit: int = SAMPLE_LIMIT) -> dict[str, Any]:
    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=True)
    ontology = yaml.safe_load(Path(ontology_path).read_text(encoding="utf-8"))
    sources = select_sources(Path(shortlist_path), 50)
    if not sources:
        raise ValueError("No usable structured resources found in the Part 1 shortlist.")
    agent = build_agent()
    manifest: list[dict[str, Any]] = []
    summary = []
    for source in sources:
        result = agent.invoke({"source": source, "ontology": ontology, "sample_limit": sample_limit})
        if result.get("error"):
            manifest.append({"source_id": source["source_id"], "error": result["error"]})
            continue
        if len(result["records"]) < sample_limit:
            manifest.append({"source_id": source["source_id"],
                             "error": f"Only {len(result['records'])} sample rows available; expected {sample_limit}"})
            continue
        if not result["mapping"]["fields"] or not result["observations"]:
            manifest.append({"source_id": source["source_id"], "error": "No confident canonical field mappings; source retained as an abstention"})
            continue
        mapping = result["mapping"]
        mapping["agent_framework"] = "LangGraph"
        mapping["model_provider"] = "OpenAI"
        mapping["model"] = result["llm_model"]
        mapping["revision_history"] = result.get("revisions", [])
        mapping["validation"] = result["validation"]
        mapping["approval"] = "pending"
        stem = re.sub(r"[^a-zA-Z0-9_-]+", "_", source["source_id"])
        for folder in ("mappings", "profiles", "observations"):
            (output / folder).mkdir(exist_ok=True)
        (output / "mappings" / f"{stem}.json").write_text(json.dumps(mapping, indent=2, ensure_ascii=False), encoding="utf-8")
        (output / "profiles" / f"{stem}.json").write_text(json.dumps(result["profile"], indent=2, ensure_ascii=False), encoding="utf-8")
        (output / "observations" / f"{stem}.jsonl").write_text("".join(json.dumps(row, ensure_ascii=False) + "\n" for row in result["observations"]), encoding="utf-8")
        manifest.append({"source_id": source["source_id"], "title": source["title"], "format": source["format"],
                         "resource_id": source["resource"].get("id"), "url": source["resource"].get("url"),
                         "licence": source.get("licence", "unknown"), "approval": "pending"})
        summary.append({"source_id": source["source_id"], "sample_rows": len(result["records"]),
                        "mapped_fields": len(mapping["fields"]), "observations": len(result["observations"]),
                        "validation_passed": result["validation"]["passed"],
                        "model": result["llm_model"], "revisions": len(result.get("revisions", [])),
                        "approval": "pending"})
        if len(summary) >= limit:
            break
    report = {"engine_version": ENGINE_VERSION, "framework": "LangGraph", "model_provider": "OpenAI",
              "sources": manifest, "results": summary,
              "errors": [item for item in manifest if "error" in item]}
    (output / "source_manifest.json").write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    if len(summary) < limit:
        raise RuntimeError(f"Processed only {len(summary)} of {limit} required sources; see {output / 'source_manifest.json'} for sampling failures.")
    return report